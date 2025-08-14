# agents/report_generator_agent.py
import pandas as pd
import numpy as np
from typing import Dict, List, Any, Optional
import logging
from datetime import datetime
from io import BytesIO
import base64

# Excel processing
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    from openpyxl.chart import BarChart, Reference
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False

class ReportGeneratorAgent:
    """
    Generates various reports and exports for FMEA data
    Implements the ReportGeneratorAgent from the class diagram
    """
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
        self.excel_exporter = ExcelExporter() if HAS_EXCEL else None
    
    def export_excel(self, fmea_data: pd.DataFrame, include_charts: bool = True, 
                    include_summary: bool = True) -> bytes:
        """Export FMEA data to Excel format with formatting and charts"""
        try:
            if not HAS_EXCEL:
                # Fallback to basic CSV if Excel libraries not available
                return fmea_data.to_csv(index=False).encode('utf-8')
            
            return self.excel_exporter.create_formatted_excel(
                fmea_data, include_charts, include_summary
            )
            
        except Exception as e:
            self.logger.error(f"Excel export failed: {str(e)}")
            # Fallback to CSV
            return fmea_data.to_csv(index=False).encode('utf-8')
    
    def generate_excel(self, fmea_data: pd.DataFrame) -> bytes:
        """Generate Excel report - interface method for coordinator"""
        return self.export_excel(fmea_data, include_charts=True, include_summary=True)
    
    def export_data(self, data: Dict) -> Dict[str, Any]:
        """Export data in specified format - interface method for coordinator"""
        try:
            fmea_data = data.get('fmea_data')
            format_type = data.get('format', 'excel')
            
            if format_type.lower() == 'excel':
                file_data = self.export_excel(fmea_data)
                return {
                    'success': True,
                    'data': file_data,
                    'filename': f"FMEA_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                    'mime_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
                }
            elif format_type.lower() == 'csv':
                csv_data = fmea_data.to_csv(index=False).encode('utf-8')
                return {
                    'success': True,
                    'data': csv_data,
                    'filename': f"FMEA_Data_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                    'mime_type': 'text/csv'
                }
            else:
                return {
                    'success': False,
                    'error': f"Unsupported format: {format_type}"
                }
                
        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }
    
    def generate_summary_report(self, fmea_data: pd.DataFrame) -> str:
        """Generate text summary report"""
        try:
            report = f"FMEA Summary Report\n"
            report += f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
            report += "="*50 + "\n\n"
            
            # Basic statistics
            report += "OVERVIEW:\n"
            report += f"Total FMEA entries: {len(fmea_data)}\n"
            
            if 'RPN' in fmea_data.columns:
                avg_rpn = fmea_data['RPN'].mean()
                max_rpn = fmea_data['RPN'].max()
                high_risk_count = len(fmea_data[fmea_data['RPN'] > 100])
                
                report += f"Average RPN: {avg_rpn:.1f}\n"
                report += f"Maximum RPN: {max_rpn}\n"
                report += f"High-risk items (RPN > 100): {high_risk_count}\n\n"
            
            # Component breakdown
            if 'Component' in fmea_data.columns:
                report += "COMPONENT BREAKDOWN:\n"
                component_counts = fmea_data['Component'].value_counts()
                for component, count in component_counts.head(5).items():
                    report += f"• {component}: {count} failure modes\n"
                report += "\n"
            
            # Top risks
            if 'RPN' in fmea_data.columns:
                report += "TOP 5 RISKS:\n"
                top_risks = fmea_data.nlargest(5, 'RPN')
                for i, (_, row) in enumerate(top_risks.iterrows(), 1):
                    component = row.get('Component', 'Unknown')
                    failure_mode = row.get('Failure Mode', 'Unknown')
                    rpn = row.get('RPN', 0)
                    report += f"{i}. {component} - {failure_mode} (RPN: {rpn})\n"
            
            return report
            
        except Exception as e:
            self.logger.error(f"Summary report generation failed: {str(e)}")
            return f"Report generation failed: {str(e)}"


class ExcelExporter:
    """Specialized Excel export functionality"""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
    
    def create_formatted_excel(self, fmea_data: pd.DataFrame, include_charts: bool = True, 
                              include_summary: bool = True) -> bytes:
        """Create a fully formatted Excel file with multiple sheets"""
        try:
            # Create workbook
            wb = openpyxl.Workbook()
            
            # Remove default sheet
            wb.remove(wb.active)
            
            # Create FMEA data sheet
            self._create_fmea_sheet(wb, fmea_data)
            
            # Create summary sheet if requested
            if include_summary:
                self._create_summary_sheet(wb, fmea_data)
            
            # Create charts sheet if requested
            if include_charts and 'RPN' in fmea_data.columns:
                self._create_charts_sheet(wb, fmea_data)
            
            # Save to bytes
            excel_buffer = BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)
            
            return excel_buffer.getvalue()
            
        except Exception as e:
            self.logger.error(f"Excel creation failed: {str(e)}")
            raise
    
    def _create_fmea_sheet(self, workbook, fmea_data: pd.DataFrame):
        """Create the main FMEA data sheet"""
        ws = workbook.create_sheet("FMEA Data")
        
        # Add title
        ws['A1'] = "Failure Mode and Effects Analysis (FMEA)"
        ws['A1'].font = Font(size=16, bold=True)
        ws['A1'].fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        ws['A1'].font = Font(color="FFFFFF", size=16, bold=True)
        
        # Add timestamp
        ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws['A2'].font = Font(italic=True)
        
        # Add data starting from row 4
        start_row = 4
        
        # Headers
        headers = list(fmea_data.columns)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # Data rows
        for row_idx, (_, row) in enumerate(fmea_data.iterrows(), start_row + 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Color code RPN values
                if headers[col_idx-1] == 'RPN' and isinstance(value, (int, float)):
                    if value > 200:
                        cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                        cell.font = Font(color="FFFFFF")
                    elif value > 100:
                        cell.fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
                    elif value > 50:
                        cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Add borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in ws.iter_rows(min_row=start_row, max_row=start_row + len(fmea_data), 
                               min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border
    
    